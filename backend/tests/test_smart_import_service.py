"""
测试 SmartImportService：自由文本 Excel 解析 + 模板生成

覆盖：
- 表头格式识别（template / freetext）
- 步骤拆分（序号行、断言尾、无序号尾行）
- target/data 提取（在X中输入Y / 点击X / 输入X）
- priority 映射（高/中/低 → P1/P2/P3）
- 缺 expected → 待补
- 模板表（9列）透传
"""
import pytest

from app.services.smart_import_service import (
    SmartImportService,
    detect_format,
    parse_free_text_row,
    split_steps,
)


# ---------- detect_format ----------

def test_detect_template_by_headers():
    headers = ["用例名", "优先级", "类型", "步骤", "动作", "目标", "数据", "预期(步)", "预期结果"]
    assert detect_format(headers) == "template"


def test_detect_freetext_by_headers():
    headers = ["用例编号", "模块", "标题", "操作步骤", "预期结果", "优先级", "测试结果", "用例截图", "失败原因"]
    assert detect_format(headers) == "freetext"


def test_detect_unknown():
    assert detect_format(["foo", "bar"]) is None


# ---------- split_steps ----------

def test_split_steps_numbered_lines():
    text = "1.登录系统\n2.点击左侧菜单采集配置\n3.点击查询"
    steps = split_steps(text)
    assert len(steps) == 3
    assert steps[0]["action"] == "登录系统"
    assert steps[2]["action"] == "点击查询"


def test_split_steps_assertion_tail_sets_expected():
    text = "1.登录系统\n2.点击查询\n断言：查询列表展示数据"
    steps = split_steps(text)
    assert len(steps) == 2
    assert steps[1]["expected"] == "查询列表展示数据"


def test_split_steps_unnumbered_line_appends_action():
    text = "1.登录系统\n进入首页\n2.点击查询"
    steps = split_steps(text)
    assert len(steps) == 2
    assert "进入首页" in steps[0]["action"]


def test_split_steps_no_assertion_expected_pending():
    text = "1.登录系统\n2.点击查询"
    steps = split_steps(text)
    assert all(s["expected"] == "待补" for s in steps)


def test_split_steps_step_numbers_sequential():
    text = "1.输入网址\n2.输入账号\n断言：登录成功"
    steps = split_steps(text)
    assert [s["step"] for s in steps] == [1, 2]


# ---------- parse_free_text_row ----------

def test_free_text_row_full_mapping():
    row = {
        "模块": "登录", "标题": "登录功能", "优先级": "P0",
        "操作步骤": "1.在登录页中输入admin\n2.点击登录按钮\n断言：提示登录成功",
        "预期结果": "登录成功",
    }
    c = parse_free_text_row(row)
    assert c["name"] == "登录-登录功能"
    assert c["priority"] == "P0"
    assert c["expected_result"] == "登录成功"
    assert c["steps"][0]["target"] == "登录页"
    assert c["steps"][0]["data"] == "admin"
    assert c["steps"][1]["target"] == "登录按钮"
    assert c["steps"][1]["expected"] == "提示登录成功"


def test_free_text_row_priority_chinese_mapping():
    assert parse_free_text_row({"标题": "x", "优先级": "高", "操作步骤": "1.点击", "预期结果": "ok"})["priority"] == "P1"
    assert parse_free_text_row({"标题": "x", "优先级": "中", "操作步骤": "1.点击", "预期结果": "ok"})["priority"] == "P2"
    assert parse_free_text_row({"标题": "x", "优先级": "低", "操作步骤": "1.点击", "预期结果": "ok"})["priority"] == "P3"
    assert parse_free_text_row({"标题": "x", "优先级": "紧急", "操作步骤": "1.点击", "预期结果": "ok"})["priority"] == "P1"


def test_free_text_row_no_module_name_is_title():
    c = parse_free_text_row({"标题": "IEMI查询验证", "操作步骤": "1.点击查询", "预期结果": "查询成功"})
    assert c["name"] == "IEMI查询验证"


def test_free_text_row_missing_expected_result_uses_first_step():
    c = parse_free_text_row({"标题": "x", "操作步骤": "1.点击查询\n断言：列表刷新", "预期结果": ""})
    assert c["expected_result"] == "列表刷新"


def test_free_text_row_expected_result_fallback_pending():
    c = parse_free_text_row({"标题": "x", "操作步骤": "1.点击查询", "预期结果": ""})
    assert c["expected_result"] == "待补"


# ---------- generate_template / parse_preview ----------

def test_generate_template_xlsx():
    data = SmartImportService.generate_template()
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == ["用例名", "优先级", "类型", "步骤", "动作", "目标", "数据", "预期(步)", "预期结果"]
    assert ws.max_row >= 2  # example row


def test_parse_preview_template_format():
    data = SmartImportService.generate_template()
    result = SmartImportService.parse_preview(data, "xlsx")
    assert result["format"] == "template"
    assert len(result["cases"]) >= 1
    case = result["cases"][0]
    assert case["name"]
    assert case["priority"] in ("P0", "P1", "P2", "P3")
    assert len(case["steps"]) >= 1


def test_parse_preview_freetext_format():
    from openpyxl import Workbook
    import io
    wb = Workbook()
    ws = wb.active
    ws.append(["用例编号", "模块", "标题", "操作步骤", "预期结果", "优先级", "测试结果", "用例截图", "失败原因"])
    ws.append(["TC001", "登录", "登录功能", "1.输入网址\n2.输入账号\n断言：登录成功", "登录成功", "P0", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    result = SmartImportService.parse_preview(buf.getvalue(), "xlsx")
    assert result["format"] == "freetext"
    assert len(result["cases"]) == 1
    c = result["cases"][0]
    assert c["name"] == "登录-登录功能"
    assert c["priority"] == "P0"
    assert len(c["steps"]) == 2
    assert c["steps"][1]["expected"] == "登录成功"
